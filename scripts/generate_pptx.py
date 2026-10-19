import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color definitions
    NAVY = RGBColor(26, 54, 93)      # #1a365d
    BLUE = RGBColor(49, 130, 206)    # #3182ce
    DARK_GRAY = RGBColor(45, 55, 72) # #2d3748
    LIGHT_BG = RGBColor(247, 250, 252) # #f7fafc
    WHITE = RGBColor(255, 255, 255)
    GREEN = RGBColor(40, 167, 69)

    def add_header(slide, title_text, category_text="PARALLEL COMPUTING ARCHITECTURE PROGRAMMING"):
        # Header banner
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.1))
        shape.fill.solid()
        shape.fill.fore_color.rgb = NAVY
        shape.line.color.rgb = NAVY

        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(11.7), Inches(0.9))
        tf = tx_box.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = BLUE

        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = WHITE

    # -------------------------------------------------------------
    # Slide 1: Title Slide
    # -------------------------------------------------------------
    slide_layout = prs.slide_layouts[6] # Blank
    s1 = prs.slides.add_slide(slide_layout)

    bg = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()

    tb = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(4.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Parallel Cloud Gaming Resource Management System"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = WHITE

    p2 = tf.add_paragraph()
    p2.text = "High-Performance Discrete-Event Scheduling Using C, OpenMP, and MPI"
    p2.font.size = Pt(22)
    p2.font.color.rgb = BLUE
    p2.space_before = Pt(12)

    p3 = tf.add_paragraph()
    p3.text = "\nCourse: Parallel Computing Architecture Programming (PCAP) Lab Mini Project\nTarget Machine: Apple Silicon Multi-Core Architecture (macOS)\nDeliverables: Sequential, OpenMP, MPI, Hybrid Backends & Research Reports"
    p3.font.size = Pt(14)
    p3.font.color.rgb = RGBColor(226, 232, 240)
    p3.space_before = Pt(20)

    # -------------------------------------------------------------
    # Slide 2: Introduction & Motivation
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(slide_layout)
    add_header(s2, "Introduction & Motivation")
    tb = s2.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.5))
    tf = tb.text_frame
    tf.word_wrap = True

    items = [
        ("Cloud Gaming Imperative:", "Renders 3D graphics on remote edge-cloud servers and streams video to thin clients. Demands strict round-trip motion-to-photon latency under 50-100 ms to preserve QoE."),
        ("Multi-Dimensional Packing:", "Incoming gaming sessions simultaneously constrain four physical hardware metrics: CPU cores, GPU rendering units, RAM capacity, and network streaming bandwidth."),
        ("The Scalability Bottleneck:", "In large-scale geo-distributed edge environments with 1,000 to 10,000+ candidate servers, sequential feasibility checks and fitness scoring create severe admission delays during peak bursts."),
        ("Parallel Computing Solution:", "Leveraging shared-memory threading (OpenMP), message-passing processes (MPI), and hybrid decomposition allows parallel exploration of candidate servers to achieve real-time response times.")
    ]
    for title, desc in items:
        p = tf.add_paragraph()
        p.text = f"•  {title} "
        p.font.bold = True
        p.font.size = Pt(15)
        p.font.color.rgb = NAVY
        p.space_before = Pt(14)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = DARK_GRAY

    # -------------------------------------------------------------
    # Slide 3: Problem Statement & Objectives
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(slide_layout)
    add_header(s3, "Problem Statement & Project Objectives")
    tb = s3.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Mathematical Problem Formulation:"
    p.font.bold = True
    p.font.size = Pt(16)
    p.font.color.rgb = NAVY

    p = tf.add_paragraph()
    p.text = "Assign each dynamic request r to an eligible server s that minimizes composite score S = 0.5 * L + 0.5 * U, subject to:"
    p.font.size = Pt(14)
    p.font.color.rgb = DARK_GRAY

    sub_items = [
        "1. Capacity Constraints: Sum of active session demands <= Total server capacity across CPU, GPU, RAM, Bandwidth.",
        "2. Network Latency Constraint: Modeled regional latency(player_region, server_region) <= Request latency limit.",
        "3. Lifecycle Conservation: Automatic, exact resource release upon session duration expiration (t_end = t_arr + duration)."
    ]
    for it in sub_items:
        p = tf.add_paragraph()
        p.text = f"    {it}"
        p.font.size = Pt(13)
        p.font.color.rgb = BLUE

    p = tf.add_paragraph()
    p.text = "\nCore Project Objectives:"
    p.font.bold = True
    p.font.size = Pt(16)
    p.font.color.rgb = NAVY

    objs = [
        "1. Reproducible C Simulator: Build reference sequential engine and discrete-event min-heap queue.",
        "2. Multi-Model Parallel Acceleration: Implement OpenMP, MPI, and Hybrid MPI+OpenMP candidate evaluation.",
        "3. 100% Cross-Backend Equivalence: Enforce identical per-request decisions and bitwise log agreement.",
        "4. Systematic Benchmarking: Evaluate speedup, efficiency, and overhead across worker budgets (1, 2, 4) up to 10,000 servers.",
        "5. Policy Trade-Off Analysis: Contrast Scored Allocation against the greedy First-Fit heuristic baseline."
    ]
    for it in objs:
        p = tf.add_paragraph()
        p.text = f"•  {it}"
        p.font.size = Pt(13)
        p.font.color.rgb = DARK_GRAY

    # -------------------------------------------------------------
    # Slide 4: Literature Review & Gap Analysis
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(slide_layout)
    add_header(s4, "Literature Review (10 Papers Surveyed: 2020-2025)")
    tb = s4.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.7), Inches(5.8))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Systematic Survey of High-Impact Publications (IEEE, ACM, Springer, Elsevier):"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = NAVY

    papers = [
        ("Deng et al. (IEEE IoT-J 2023):", "Multi-Agent DRL for edge cloud gaming.", "High QoE adaptability, but high offline training and inference delay."),
        ("Zhao et al. (IEEE TPDS 2021):", "Multi-GPU game co-location.", "Eliminates frame rate degradation via multi-dimensional bin packing."),
        ("Hossain et al. (IEEE TNSM 2022):", "MEC latency-aware offloading.", "Analyzed delay trade-offs; evaluated on homogeneous sequential models."),
        ("Tan et al. (IEEE/ACM ToN 2021):", "Lyapunov online admission control.", "Theoretical queue bounds, but computationally unviable at 10,000 servers."),
        ("Mishra & Dash (Springer 2022):", "OpenMP vs. MPI on NUMA systems.", "Quantified communication overheads in reduction-heavy algorithms."),
        ("Gomez-Martin (Elsevier JSA 2022):", "Hybrid MPI+OpenMP heuristics.", "Hierarchical decomposition reduces collective network message traffic.")
    ]
    for auth, topic, insight in papers:
        p = tf.add_paragraph()
        p.text = f"•  {auth} "
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = NAVY
        r1 = p.add_run()
        r1.text = f"{topic} — "
        r1.font.bold = False
        r1.font.color.rgb = BLUE
        r2 = p.add_run()
        r2.text = insight
        r2.font.bold = False
        r2.font.color.rgb = DARK_GRAY

    p = tf.add_paragraph()
    p.text = "\nIdentified Research Gap: Existing schedulers either rely on slow sequential heuristics or black-box offline ML models. A high-performance, deterministic C simulator evaluating OpenMP vs. MPI vs. Hybrid scaling directly on multi-core unified memory architectures has been unaddressed."
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = RGBColor(197, 48, 48)

    # -------------------------------------------------------------
    # Slide 5: System Architecture & Workflow Diagram
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(slide_layout)
    add_header(s5, "System Architecture & Simulation Workflow")

    flowchart_path = "plots/architecture_flowchart.png"
    if os.path.exists(flowchart_path):
        s5.shapes.add_picture(flowchart_path, Inches(0.8), Inches(1.3), height=Inches(5.8))

    tb = s5.shapes.add_textbox(Inches(5.8), Inches(1.3), Inches(6.8), Inches(5.8))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Key Architectural Modules:"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = NAVY

    arch_pts = [
        ("Input/Workload Module (src/input.c):", "Reads servers, requests, and 4x4 latency matrices with high-speed CSV parsing and monotonic timer."),
        ("Event Priority Min-Heap (src/events.c):", "Tracks active sessions ordered by end_time. Drains expired sessions before each arrival to restore server capacities."),
        ("Scoring Kernel (src/scoring.c):", "Deterministic feasibility checks across 4 resources and score calculation: S = 0.5 * L + 0.5 * U."),
        ("Four Backends (src/backend_*.c):", "1. Sequential: Reference loop scan.\n2. OpenMP: Thread-private candidate reduction.\n3. MPI: Contiguous partition ownership + MPI_Allreduce.\n4. Hybrid: MPI domain split + OpenMP thread scan.")
    ]
    for t, d in arch_pts:
        p = tf.add_paragraph()
        p.text = f"• {t} "
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = BLUE
        p.space_before = Pt(8)
        run = p.add_run()
        run.text = d
        run.font.bold = False
        run.font.color.rgb = DARK_GRAY

    # -------------------------------------------------------------
    # Slide 6: Mathematical Model & Algorithms
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(slide_layout)
    add_header(s6, "Mathematical Model & Algorithms")
    tb = s6.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.7), Inches(5.8))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "1. Feasibility Condition (Eligibility Filter):"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = NAVY

    p = tf.add_paragraph()
    p.text = "    Latency(region_r, region_s) <= Latency_Limit(r)\n    AND Server_Avail[res] >= Request_Demand[res]  FOR ALL res in {CPU, GPU, RAM, Bandwidth}"
    p.font.size = Pt(13)
    p.font.color.rgb = BLUE

    p = tf.add_paragraph()
    p.text = "\n2. Normalized Composite Score Formulation:"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = NAVY

    p = tf.add_paragraph()
    p.text = "    Normalized Latency: L = Latency / Latency_Limit\n    Resource Fractions: u[res] = (Total[res] - Avail[res] + Demand[res]) / Total[res]\n    Bottleneck Utilization: U = MAX(u[CPU], u[GPU], u[RAM], u[Bandwidth])\n    Score = 0.5 * L + 0.5 * U"
    p.font.size = Pt(13)
    p.font.color.rgb = BLUE

    p = tf.add_paragraph()
    p.text = "\n3. Deterministic Winner Comparison Rule:"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = NAVY

    p = tf.add_paragraph()
    p.text = "    Lowest score wins. Ties broken by lowest Global Server ID.\n    Candidate representation: struct { double score; int server_id; }\n    Directly compatible with MPI_Allreduce(..., MPI_DOUBLE_INT, MPI_MINLOC, MPI_COMM_WORLD)."
    p.font.size = Pt(13)
    p.font.color.rgb = DARK_GRAY

    # -------------------------------------------------------------
    # Slide 7: Implementation Details of the 4 Backends
    # -------------------------------------------------------------
    s7 = prs.slides.add_slide(slide_layout)
    add_header(s7, "Implementation Architecture Across 4 Backends")
    tb = s7.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.7), Inches(5.8))
    tf = tb.text_frame
    tf.word_wrap = True

    b_details = [
        ("Sequential C (src/backend_seq.c):", "Single-threaded scan through all N servers. Serves as the correctness oracle and baseline for speedup calculation. Supports both Scored and First-Fit policies."),
        ("OpenMP Shared-Memory (src/backend_omp.c):", "Uses #pragma omp parallel with #pragma omp for schedule(static). Each thread evaluates assigned servers into a private candidate variable without locks. Winners merged once per thread. Single-threaded master commits reservation."),
        ("MPI Message-Passing (src/backend_mpi.c):", "Partitions N servers contiguously across P ranks: count = N/P + (rank < N%P ? 1 : 0). Rank 0 broadcasts request (MPI_Bcast). Each rank scans local partition. Global winner reduced via MPI_Allreduce(MPI_MINLOC). Only owning rank reserves capacity."),
        ("Hybrid MPI+OpenMP (src/backend_hybrid.c):", "Initialized via MPI_Init_thread(MPI_THREAD_FUNNELED). Combines MPI domain decomposition across ranks with OpenMP multi-threaded scanning inside each rank. Evaluated at fixed worker budget W = P x T.")
    ]
    for b_title, b_desc in b_details:
        p = tf.add_paragraph()
        p.text = f"• {b_title} "
        p.font.bold = True
        p.font.size = Pt(14)
        p.font.color.rgb = NAVY
        p.space_before = Pt(10)
        run = p.add_run()
        run.text = b_desc
        run.font.bold = False
        run.font.color.rgb = DARK_GRAY

    # -------------------------------------------------------------
    # Slide 8: Experimental Setup & Hardware Specifications
    # -------------------------------------------------------------
    s8 = prs.slides.add_slide(slide_layout)
    add_header(s8, "Experimental Setup & Target Hardware")
    tb = s8.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.7), Inches(5.5))
    tf = tb.text_frame
    tf.word_wrap = True

    hw_info = [
        ("Host Machine:", "Apple Silicon MacBook Air (Apple M-series architecture)"),
        ("CPU Configuration:", "10 logical cores (high-performance + high-efficiency cores)"),
        ("Unified Memory:", "16 GB unified RAM (17,179,869,184 bytes); 512 GB SSD storage"),
        ("Compiler Toolchain:", "Homebrew LLVM Clang version 23.1.3 (-std=c11 -O2 -Wall -Wextra)"),
        ("OpenMP Runtime:", "Homebrew libomp version 23.1.3 (-fopenmp)"),
        ("MPI Environment:", "Open MPI version 5.0.11 (native arm64 build; mpicc / mpirun)"),
        ("Python Environment:", "Python 3.12 (.venv) with Matplotlib 3.11.2, Pandas 3.0.6, NumPy 2.5.4"),
        ("Benchmark Workloads:", "100, 1,000, and 10,000 server pools evaluated under 1,000 high-arrival requests across 5 repetitions")
    ]
    for label, val in hw_info:
        p = tf.add_paragraph()
        p.text = f"•  {label} "
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = NAVY
        p.space_before = Pt(6)
        run = p.add_run()
        run.text = val
        run.font.bold = False
        run.font.color.rgb = DARK_GRAY

    # -------------------------------------------------------------
    # Slide 9: Correctness Gates Verification (12 Gates)
    # -------------------------------------------------------------
    s9 = prs.slides.add_slide(slide_layout)
    add_header(s9, "Correctness Gates & Cross-Backend Equivalence")
    tb = s9.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.7), Inches(5.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "12 Rigorous Correctness Gates Verified via Automated Test Suite (tests/test_correctness.c):"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = NAVY

    gates = [
        ("Gate 1 & 2: Boundary Feasibility", "Exact fit allocates completely (0 capacity remaining); resource shortage rejects without state mutation."),
        ("Gate 3: Latency Invariant", "Rejects candidate servers exceeding request latency limit even if resources are abundant."),
        ("Gate 4: Deterministic Tie-Breaking", "Equal score candidates break ties deterministically by selecting lowest Global Server ID."),
        ("Gate 5: Event Ordering", "Session completions occurring at current arrival timestamp are released before next allocation."),
        ("Gate 6 & 7: Lifecycle Conservation", "Final event drain restores 100.0% original capacities; partition math handles uneven N/P."),
        ("Gate 8-12: Bitwise Backend Equivalence", "Sequential == OpenMP == MPI == Hybrid: 100% identical decision logs across all requests.")
    ]
    for g_title, g_desc in gates:
        p = tf.add_paragraph()
        p.text = f"[PASS]  {g_title}: "
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = GREEN
        p.space_before = Pt(8)
        run = p.add_run()
        run.text = g_desc
        run.font.bold = False
        run.font.color.rgb = DARK_GRAY

    # -------------------------------------------------------------
    # Slide 10: Benchmark Results & Scalability Figures
    # -------------------------------------------------------------
    s10 = prs.slides.add_slide(slide_layout)
    add_header(s10, "Empirical Benchmark Results: Runtime & Speedup")

    p1 = "plots/fig1_runtime_vs_workers.png"
    p2 = "plots/fig2_speedup_efficiency.png"

    if os.path.exists(p1):
        s10.shapes.add_picture(p1, Inches(0.8), Inches(1.4), width=Inches(5.7))
    if os.path.exists(p2):
        s10.shapes.add_picture(p2, Inches(6.8), Inches(1.4), width=Inches(5.7))

    tb = s10.shapes.add_textbox(Inches(0.8), Inches(5.8), Inches(11.7), Inches(1.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Key Takeaway: At 10,000 servers, MPI achieves 2.92x speedup on 4 ranks (core time: 0.0125s vs 0.0366s sequential), successfully amortizing collective broadcast and reduction overheads. OpenMP runtime remains constant due to thread management overhead."
    p.font.size = Pt(12)
    p.font.color.rgb = NAVY
    p.font.bold = True

    # -------------------------------------------------------------
    # Slide 11: Scaling Behavior & Crossover Points
    # -------------------------------------------------------------
    s11 = prs.slides.add_slide(slide_layout)
    add_header(s11, "Scaling Behavior Across Server Pool Scales")

    p3 = "plots/fig3_scaling_servers.png"
    if os.path.exists(p3):
        s11.shapes.add_picture(p3, Inches(0.8), Inches(1.4), width=Inches(6.2))

    tb = s11.shapes.add_textbox(Inches(7.3), Inches(1.4), Inches(5.2), Inches(5.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Runtime Scaling & Crossover Insights:"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = NAVY

    pts = [
        ("Small Scale (100 Servers):", "Sequential C is fastest (0.0004s). OpenMP thread synchronization and MPI IPC socket costs outweigh parallel compute savings."),
        ("Medium Scale (1,000 Servers):", "Crossover region: MPI 4 ranks (0.0027s) outperforms Sequential (0.0037s) as search computational density increases."),
        ("Large Scale (10,000 Servers):", "MPI achieves dominant performance: 0.0125s vs 0.0366s (2.92x speedup). Collective communication overhead is fully amortized."),
        ("Hybrid Split:", "Hybrid 2 ranks x 2 threads achieves 0.0256s, offering a viable middle ground between thread synchronization and IPC communication.")
    ]
    for t, d in pts:
        p = tf.add_paragraph()
        p.text = f"• {t} "
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = BLUE
        p.space_before = Pt(8)
        run = p.add_run()
        run.text = d
        run.font.bold = False
        run.font.color.rgb = DARK_GRAY

    # -------------------------------------------------------------
    # Slide 12: Policy Evaluation: Scored vs First-Fit
    # -------------------------------------------------------------
    s12 = prs.slides.add_slide(slide_layout)
    add_header(s12, "Policy Comparison: Scored Allocation vs. First-Fit")

    p4 = "plots/fig4_policy_comparison.png"
    if os.path.exists(p4):
        s12.shapes.add_picture(p4, Inches(0.8), Inches(1.4), width=Inches(6.2))

    tb = s12.shapes.add_textbox(Inches(7.3), Inches(1.4), Inches(5.2), Inches(5.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "QoS and Latency Policy Comparison:"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = NAVY

    pol_pts = [
        ("54.7% Latency Reduction:", "Scored allocation achieves an average network delay of 14.22 ms compared to 31.43 ms under First-Fit heuristic."),
        ("Balanced Server Utilization:", "Scored policy evaluates candidate resource headroom, preventing server hot-spotting and premature GPU exhaustion."),
        ("The First-Fit Trade-off:", "First-Fit terminates search early at the first feasible server (0.007s runtime), but incurs over 2x worse gaming latency."),
        ("Academic Justification:", "Demonstrates that parallel acceleration is crucial: it makes high-quality Scored allocation fast enough for real-time gaming.")
    ]
    for t, d in pol_pts:
        p = tf.add_paragraph()
        p.text = f"• {t} "
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = GREEN
        p.space_before = Pt(8)
        run = p.add_run()
        run.text = d
        run.font.bold = False
        run.font.color.rgb = DARK_GRAY

    # -------------------------------------------------------------
    # Slide 13: Live Demonstration Guide
    # -------------------------------------------------------------
    s13 = prs.slides.add_slide(slide_layout)
    add_header(s13, "Project Demonstration & Execution Guide")
    tb = s13.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.7), Inches(5.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Commands for Live Evaluation & Demonstration:"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = NAVY

    cmds = [
        ("1. Automated Correctness Suite (make test):", "./test_correctness\nExecutes all 12 correctness gates; verifies boundary conditions and 100% backend equivalence."),
        ("2. Sequential Reference Run:", "./sequential --servers data/servers_1k.csv --requests data/requests_heavy_1k.csv --policy scored --log results/log_seq.txt"),
        ("3. OpenMP Multi-Threaded Run:", "./openmp --servers data/servers_1k.csv --requests data/requests_heavy_1k.csv --threads 4 --policy scored --log results/log_omp.txt"),
        ("4. MPI Distributed-Memory Run:", "mpirun -np 4 ./mpi --servers data/servers_1k.csv --requests data/requests_heavy_1k.csv --policy scored --log results/log_mpi.txt"),
        ("5. Hybrid MPI+OpenMP Run:", "mpirun -np 2 ./hybrid --servers data/servers_1k.csv --requests data/requests_heavy_1k.csv --threads 2 --policy scored --log results/log_hybrid.txt")
    ]
    for title, cmd in cmds:
        p = tf.add_paragraph()
        p.text = f"• {title}"
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = BLUE
        p.space_before = Pt(6)
        p_c = tf.add_paragraph()
        p_c.text = f"   {cmd}"
        p_c.font.size = Pt(11)
        p_c.font.color.rgb = DARK_GRAY

    # -------------------------------------------------------------
    # Slide 14: Team Contributions & Role Distribution
    # -------------------------------------------------------------
    s14 = prs.slides.add_slide(slide_layout)
    add_header(s14, "Team Contributions & Assessment Alignment")
    tb = s14.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.7), Inches(5.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Role Distribution Across Project Lifecycle (Mini Project Guideline Requirement):"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = NAVY

    roles = [
        ("Architecture & Core Discrete-Event Engine:", "Design of data models (models.h), priority min-heap event queue (events.c), session lifecycle restoration, and input CSV parsers."),
        ("Sequential & OpenMP Implementation:", "Development of reference sequential scheduler, OpenMP static loop distribution, thread-private candidate reduction, and correctness gates test suite."),
        ("MPI & Hybrid Distributed Design:", "Implementation of contiguous server partition math, MPI_Bcast request distribution, collective MPI_Allreduce(MPI_MINLOC) reduction, and MPI_THREAD_FUNNELED hybrid integration."),
        ("Benchmarking, Graphing & Report Authoring:", "Execution of 135 benchmark runs, Matplotlib plotting script, IEEE & Springer conference reports, and presentation slide synthesis.")
    ]
    for r_title, r_desc in roles:
        p = tf.add_paragraph()
        p.text = f"•  {r_title} "
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = NAVY
        p.space_before = Pt(10)
        run = p.add_run()
        run.text = r_desc
        run.font.bold = False
        run.font.color.rgb = DARK_GRAY

    # -------------------------------------------------------------
    # Slide 15: Conclusion & Future Scope
    # -------------------------------------------------------------
    s15 = prs.slides.add_slide(slide_layout)
    add_header(s15, "Conclusion & Future Work")
    tb = s15.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.7), Inches(5.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Project Summary & Key Milestones Achieved:"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = NAVY

    concls = [
        ("Full 3+1 Backend Simulator Complete:", "Built Sequential, OpenMP, MPI, and Hybrid engines in native C with zero external runtime dependencies."),
        ("2.92x Empirical Speedup:", "Demonstrated that distributed-partition MPI scales effectively for 10,000 servers on Apple Silicon multi-core hardware."),
        ("54.7% Latency Advantage:", "Proved that Scored allocation significantly outperforms First-Fit heuristics in user QoE network latency."),
        ("100% Cross-Backend Equivalence:", "Passed all 12 correctness gates with identical decision logs across all modes."),
        ("Academic Deliverables Prepared:", "Full literature survey (10 papers), IEEE conference draft, and Springer Nature conference draft completed.")
    ]
    for t, d in concls:
        p = tf.add_paragraph()
        p.text = f"•  {t} "
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = BLUE
        p.space_before = Pt(6)
        run = p.add_run()
        run.text = d
        run.font.bold = False
        run.font.color.rgb = DARK_GRAY

    p = tf.add_paragraph()
    p.text = "\nFuture Work: 1) Physical multi-node cluster benchmarking across InfiniBand; 2) Request batching to reduce MPI collective frequency; 3) GPU acceleration with CUDA/OpenACC."
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = NAVY

    # Save presentation
    prs.save("presentation.pptx")
    print("Successfully generated presentation.pptx (15 slides)")

if __name__ == '__main__':
    create_deck()
